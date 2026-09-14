import os
import json
import re
import base64
from datetime import datetime, timezone
from urllib.parse import quote
from urllib.request import Request, urlopen
import xml.etree.ElementTree as ET

from fastapi import FastAPI, UploadFile, File
from fastapi.responses import HTMLResponse
from pydantic import BaseModel
from huggingface_hub import InferenceClient
from pypdf import PdfReader
from docx import Document

app = FastAPI()

HF_TOKEN = os.getenv("HF_TOKEN")
MODEL = "zai-org/GLM-5.3-Flash"

client = InferenceClient(
    provider="novita",
    api_key=HF_TOKEN
)

VISION_MODEL = "Qwen/Qwen3-VL-30B-A3B-Instruct"

vision_client = InferenceClient(
    provider="novita",
    api_key=HF_TOKEN
)


class ChatRequest(BaseModel):
    message: str
    history: list = []
    personality: str = "Friendly"
    response_style: str = "Balanced"
    custom_instructions: str = ""
    document_text: str = ""



class FileRequest(BaseModel):
    message: str
    document_text: str
    history: list = []
    personality: str = "Friendly"
    response_style: str = "Balanced"
    custom_instructions: str = ""


PERSONALITIES = {
    "Friendly": "Be friendly, natural, helpful and easy to understand.",
    "Teacher": "Explain clearly like a good teacher using simple examples.",
    "Coding Assistant": "Focus on accurate programming help and practical debugging.",
    "Professional": "Use a professional, clear and structured communication style."
}

STYLES = {
    "Short": "Keep answers concise and direct.",
    "Balanced": "Give balanced answers with useful explanation.",
    "Detailed": "Give detailed explanations and examples when useful."
}



def extract_document_text(filename, raw_bytes):
    from io import BytesIO
    lower_name = filename.lower()

    if lower_name.endswith(".pdf"):
        reader = PdfReader(BytesIO(raw_bytes))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()

    if lower_name.endswith(".docx"):
        document = Document(BytesIO(raw_bytes))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts).strip()

    if lower_name.endswith(".txt"):
        return raw_bytes.decode("utf-8", errors="replace").strip()

    raise ValueError("Unsupported file type. Use PDF, DOCX, or TXT.")


def limit_document_text(text, max_chars=50000):
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n[Document truncated for processing.]"


def needs_live_search(message):
    keywords = [
        "latest", "today", "current", "recent", "news",
        "right now", "this week", "this month",
        "search", "trending", "weather", "temperature", "forecast"
    ]
    text = message.lower()
    return any(word in text for word in keywords)


def google_news_search(query):
    try:
        url = (
            "https://news.google.com/rss/search?q="
            + quote(query)
            + "&hl=en-IN&gl=IN&ceid=IN:en"
        )
        request = Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urlopen(request, timeout=10) as response:
            data = response.read().decode("utf-8")

        root = ET.fromstring(data)
        results = []

        for item in root.findall(".//item")[:5]:
            title = item.findtext("title") or ""
            date = item.findtext("pubDate") or ""
            link = item.findtext("link") or ""
            source = item.findtext("source") or "Google News"
            if title:
                results.append({
                    "title": title,
                    "date": date,
                    "link": link,
                    "source": source
                })

        return results
    except Exception as error:
        print("NEWS ERROR:", error)
        return []


def wikipedia_search(query):
    try:
        url = (
            "https://en.wikipedia.org/w/api.php?"
            "action=query&format=json&list=search"
            "&srsearch=" + quote(query) + "&srlimit=3"
        )
        request = Request(url, headers={"User-Agent": "RAIZEN-AI/1.0"})

        with urlopen(request, timeout=10) as response:
            data = json.loads(response.read().decode("utf-8"))

        results = []
        for item in data.get("query", {}).get("search", []):
            title = item.get("title", "")
            results.append({
                "title": title,
                "snippet": re.sub("<.*?>", "", item.get("snippet", "")),
                "link": "https://en.wikipedia.org/wiki/" + quote(title.replace(" ", "_")),
                "source": "Wikipedia"
            })

        return results
    except Exception as error:
        print("WIKIPEDIA ERROR:", error)
        return []


# Direct coordinates prevent Cuttack from being incorrectly resolved
# to a nearby locality such as Urali.
CITY_COORDINATES = {
    "cuttack": {
        "latitude": 20.4625,
        "longitude": 85.8830,
        "name": "Cuttack",
        "country": "India"
    },
    "bhubaneswar": {
        "latitude": 20.2961,
        "longitude": 85.8245,
        "name": "Bhubaneswar",
        "country": "India"
    },
    "delhi": {
        "latitude": 28.6139,
        "longitude": 77.2090,
        "name": "Delhi",
        "country": "India"
    },
    "mumbai": {
        "latitude": 19.0760,
        "longitude": 72.8777,
        "name": "Mumbai",
        "country": "India"
    },
    "kolkata": {
        "latitude": 22.5726,
        "longitude": 88.3639,
        "name": "Kolkata",
        "country": "India"
    },
    "chennai": {
        "latitude": 13.0827,
        "longitude": 80.2707,
        "name": "Chennai",
        "country": "India"
    },
    "hyderabad": {
        "latitude": 17.3850,
        "longitude": 78.4867,
        "name": "Hyderabad",
        "country": "India"
    },
    "bangalore": {
        "latitude": 12.9716,
        "longitude": 77.5946,
        "name": "Bengaluru",
        "country": "India"
    },
    "bengaluru": {
        "latitude": 12.9716,
        "longitude": 77.5946,
        "name": "Bengaluru",
        "country": "India"
    }
}


def weather_condition(code):
    codes = {
        0: "Clear sky",
        1: "Mainly clear",
        2: "Partly cloudy",
        3: "Overcast",
        45: "Fog",
        48: "Fog",
        51: "Light drizzle",
        53: "Drizzle",
        55: "Heavy drizzle",
        56: "Light freezing drizzle",
        57: "Heavy freezing drizzle",
        61: "Light rain",
        63: "Moderate rain",
        65: "Heavy rain",
        66: "Light freezing rain",
        67: "Heavy freezing rain",
        71: "Light snow",
        73: "Moderate snow",
        75: "Heavy snow",
        77: "Snow grains",
        80: "Rain showers",
        81: "Rain showers",
        82: "Heavy rain showers",
        85: "Snow showers",
        86: "Heavy snow showers",
        95: "Thunderstorm",
        96: "Thunderstorm with hail",
        99: "Thunderstorm with heavy hail"
    }
    return codes.get(code, "Unknown")


def get_city_location(city):
    city = city.strip()
    if not city:
        return None

    known = CITY_COORDINATES.get(city.lower())
    if known:
        return known

    geo_url = (
        "https://geocoding-api.open-meteo.com/v1/search"
        "?name=" + quote(city)
        + "&count=10&language=en&format=json"
    )
    request = Request(
        geo_url,
        headers={"User-Agent": "RAIZEN-AI/1.0"}
    )

    with urlopen(request, timeout=10) as response:
        geo_data = json.loads(response.read().decode("utf-8"))

    locations = geo_data.get("results", [])
    if not locations:
        return None

    city_clean = city.lower()
    for item in locations:
        if item.get("name", "").strip().lower() == city_clean:
            return item

    return locations[0]


def weather_open_meteo(city):
    try:
        location = get_city_location(city)
        if not location:
            return None

        latitude = location.get("latitude")
        longitude = location.get("longitude")

        if latitude is None or longitude is None:
            return None

        weather_url = (
            "https://api.open-meteo.com/v1/forecast?"
            "latitude=" + str(latitude)
            + "&longitude=" + str(longitude)
            + "&current=temperature_2m,"
            "relative_humidity_2m,"
            "apparent_temperature,"
            "precipitation,"
            "wind_speed_10m,"
            "weather_code"
            "&timezone=auto"
        )

        request = Request(
            weather_url,
            headers={"User-Agent": "RAIZEN-AI/1.0"}
        )

        with urlopen(request, timeout=10) as response:
            weather_data = json.loads(response.read().decode("utf-8"))

        current = weather_data.get("current", {})
        condition = weather_condition(current.get("weather_code"))

        return (
            "Current weather for "
            + location.get("name", city)
            + ", "
            + location.get("country", "")
            + "\n\n"
            + "Condition: " + condition
            + "\n"
            + "Temperature: "
            + str(current.get("temperature_2m"))
            + "Â°C\n"
            + "Feels like: "
            + str(current.get("apparent_temperature"))
            + "Â°C\n"
            + "Humidity: "
            + str(current.get("relative_humidity_2m"))
            + "%\n"
            + "Precipitation: "
            + str(current.get("precipitation"))
            + " mm\n"
            + "Wind speed: "
            + str(current.get("wind_speed_10m"))
            + " km/h\n\n"
            + "Source: Open-Meteo"
        )
    except Exception as error:
        print("OPEN-METEO WEATHER ERROR:", error)
        return None


def weather_met_norway(city):
    try:
        location = get_city_location(city)
        if not location:
            return None

        latitude = location.get("latitude")
        longitude = location.get("longitude")

        if latitude is None or longitude is None:
            return None

        url = (
            "https://api.met.no/weatherapi/locationforecast/2.0/compact?"
            "lat=" + str(round(float(latitude), 4))
            + "&lon=" + str(round(float(longitude), 4))
        )

        request = Request(
            url,
            headers={
                "User-Agent": (
                    "RAIZEN-AI/1.0 "
                    "(https://github.com/raihankausark-beep/RAIZEN_AI)"
                )
            }
        )

        with urlopen(request, timeout=15) as response:
            data = json.loads(response.read().decode("utf-8"))

        timeseries = (
            data.get("properties", {}).get("timeseries", [])
        )
        if not timeseries:
            return None

        now = datetime.now(timezone.utc)
        selected = timeseries[0]

        for item in timeseries:
            timestamp = item.get("time")
            if not timestamp:
                continue
            try:
                item_time = datetime.fromisoformat(
                    timestamp.replace("Z", "+00:00")
                )
                if item_time >= now:
                    selected = item
                    break
            except Exception:
                continue

        data_block = selected.get("data", {})
        instant = (
            data_block.get("instant", {}).get("details", {})
        )
        next_hour = data_block.get("next_1_hours", {})
        summary = next_hour.get("summary", {})
        details = next_hour.get("details", {})

        temperature = instant.get("air_temperature")
        humidity = instant.get("relative_humidity")
        wind_speed = instant.get("wind_speed")
        precipitation = details.get("precipitation_amount")

        symbol = summary.get("symbol_code", "unknown")
        condition = (
            symbol.replace("_", " ")
            .replace("day", "")
            .replace("night", "")
            .strip()
            .title()
        )

        return (
            "Current weather for "
            + location.get("name", city)
            + ", "
            + location.get("country", "")
            + "\n\n"
            + "Condition: " + condition
            + "\n"
            + "Temperature: " + str(temperature) + "Â°C\n"
            + "Humidity: " + str(humidity) + "%\n"
            + "Precipitation: " + str(precipitation) + " mm\n"
            + "Wind speed: " + str(wind_speed) + " m/s\n\n"
            + "Source: MET Norway"
        )
    except Exception as error:
        print("MET NORWAY WEATHER ERROR:", error)
        return None


def weather_search(city):
    result = weather_open_meteo(city)
    if result:
        return result

    result = weather_met_norway(city)
    if result:
        return result

    return None


def extract_city(message):
    patterns = [
        r"weather\s+(?:in|at|of)\s+(.+)",
        r"temperature\s+(?:in|at|of)\s+(.+)",
        r"forecast\s+(?:in|at|of)\s+(.+)",
        r"weather\s+(.+)",
        r"temperature\s+(.+)"
    ]

    for pattern in patterns:
        match = re.search(
            pattern,
            message.strip(),
            re.IGNORECASE
        )

        if match:
            city = match.group(1).strip()
            city = re.sub(
                r"\b(today|now|currently|right now|tomorrow)\b",
                "",
                city,
                flags=re.IGNORECASE
            )
            city = city.strip(" ?.,!")

            if city:
                return city

    return None


def live_search(message):
    text = message.lower().strip()

    if (
        "weather" in text
        or "temperature" in text
        or "forecast" in text
    ):
        city = extract_city(message)

        if not city:
            return "WEATHER_ERROR: Please provide a city name."

        result = weather_search(city)

        if result:
            return result

        return (
            "WEATHER_ERROR: "
            "Weather data could not be retrieved right now."
        )

    news = google_news_search(message)

    if news:
        output = "LIVE NEWS RESULTS:\n\n"
        for index, item in enumerate(news, 1):
            output += (
                str(index)
                + ". "
                + item["title"]
                + "\nSource: "
                + item.get("source", "Google News")
                + "\nDate: "
                + item.get("date", "")
                + "\nLink: "
                + item.get("link", "")
                + "\n\n"
            )
        return output

    wiki = wikipedia_search(message)

    if wiki:
        output = "WEB INFORMATION:\n\n"
        for index, item in enumerate(wiki, 1):
            output += (
                str(index)
                + ". "
                + item["title"]
                + "\n"
                + item["snippet"]
                + "\nSource: "
                + item.get("source", "Wikipedia")
                + "\nLink: "
                + item.get("link", "")
                + "\n\n"
            )
        return output

    return "LIVE_SEARCH_ERROR: Live search could not retrieve data."


HTML = r"""
<!DOCTYPE html>
<html>
<head>
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>RAIZEN AI</title>
<style>
*{box-sizing:border-box}html,body{margin:0;padding:0;width:100%;min-height:100%}body{font-family:Arial,Helvetica,sans-serif;color:#f8fafc;background:radial-gradient(circle at 15% 5%,rgba(59,130,246,.22),transparent 28%),radial-gradient(circle at 85% 15%,rgba(139,92,246,.20),transparent 30%),radial-gradient(circle at 50% 100%,rgba(14,165,233,.10),transparent 35%),#05070d}.container{width:min(1120px,100%);min-height:100vh;margin:auto;padding:18px;display:flex;flex-direction:column}.header{padding:17px 20px;background:rgba(15,23,42,.72);backdrop-filter:blur(20px);-webkit-backdrop-filter:blur(20px);border:1px solid rgba(255,255,255,.09);border-radius:22px;margin-bottom:14px;display:flex;justify-content:space-between;align-items:center;box-shadow:0 14px 50px rgba(0,0,0,.28);position:sticky;top:10px;z-index:10}.logo{font-size:28px;font-weight:850;letter-spacing:.6px}.status{font-size:12px;color:#86efac;font-weight:800;letter-spacing:.5px;display:flex;align-items:center;gap:7px}.status:before{content:"";width:8px;height:8px;border-radius:50%;background:#4ade80;box-shadow:0 0 12px rgba(74,222,128,.8)}.chat{flex:1;overflow-y:auto;padding:8px 4px 22px;scroll-behavior:smooth}#welcome{text-align:center;margin:42px auto 28px;max-width:760px;padding:34px 25px;background:linear-gradient(145deg,rgba(30,41,59,.72),rgba(15,23,42,.45));border:1px solid rgba(255,255,255,.08);border-radius:28px;box-shadow:0 20px 70px rgba(0,0,0,.28);font-size:17px;line-height:1.75}#welcome::first-line{font-size:29px;font-weight:850}.message{max-width:82%;padding:14px 17px;margin:9px 0;border-radius:19px;white-space:pre-wrap;line-height:1.58;animation:messageIn .22s ease;box-shadow:0 8px 28px rgba(0,0,0,.12)}@keyframes messageIn{from{opacity:0;transform:translateY(7px)}to{opacity:1;transform:translateY(0)}}.user{margin-left:auto;background:linear-gradient(135deg,#2563eb,#7c3aed);border:1px solid rgba(255,255,255,.08)}.assistant{background:rgba(15,23,42,.86);border:1px solid rgba(255,255,255,.075)}.thinking{opacity:.72}.controls{display:flex;gap:8px;margin-bottom:9px;flex-wrap:wrap}select,.control{background:rgba(15,23,42,.88);color:#f8fafc;border:1px solid rgba(255,255,255,.10);border-radius:13px;padding:9px 11px;outline:none}.control{cursor:pointer;transition:transform .18s ease,border-color .18s ease,background .18s ease}.control:hover{transform:translateY(-1px);border-color:rgba(129,140,248,.55);background:rgba(30,41,59,.95)}.file-panel,.vision-panel{display:flex;gap:9px;align-items:center;margin-bottom:9px;flex-wrap:wrap;padding:10px 12px;background:rgba(15,23,42,.55);border:1px solid rgba(255,255,255,.065);border-radius:15px}.file-input,.vision-input{max-width:100%;color:#cbd5e1;font-size:13px}.file-name,.vision-name{font-size:12px;opacity:.76}.input-area{display:flex;gap:9px;padding-top:3px}#messageInput{flex:1;min-width:0;padding:15px 17px;background:rgba(15,23,42,.94);color:white;border:1px solid rgba(255,255,255,.11);border-radius:17px;outline:none;font-size:15px;box-shadow:0 10px 35px rgba(0,0,0,.16)}#messageInput::placeholder{color:#94a3b8}#messageInput:focus{border-color:#6366f1;box-shadow:0 0 0 3px rgba(99,102,241,.13),0 10px 35px rgba(0,0,0,.18)}#sendButton{min-width:82px;padding:0 21px;background:linear-gradient(135deg,#2563eb,#7c3aed);color:white;border:0;border-radius:17px;cursor:pointer;font-weight:800;font-size:14px;transition:transform .18s ease,filter .18s ease;box-shadow:0 10px 30px rgba(79,70,229,.24)}#sendButton:hover{transform:translateY(-1px);filter:brightness(1.08)}#sendButton:disabled{opacity:.55;cursor:not-allowed;transform:none}.message a{color:#93c5fd;text-decoration:underline;word-break:break-all}.search-badge{display:inline-block;font-size:11px;padding:3px 7px;border:1px solid #3b4d73;border-radius:999px;opacity:.8;margin-bottom:5px}.small{text-align:center;opacity:.48;font-size:11px;margin-top:11px;padding-bottom:3px}@media(max-width:650px){.container{padding:9px}.header{padding:14px 15px;border-radius:18px;top:5px}.logo{font-size:23px}.status{font-size:10px}#welcome{margin:24px auto 20px;padding:27px 17px;border-radius:23px;font-size:14px}#welcome::first-line{font-size:23px}.message{max-width:93%;font-size:14px;padding:12px 14px}.controls{gap:6px}select,.control{font-size:12px;padding:8px 9px}.input-area{position:sticky;bottom:0;padding:8px 0;background:#05070d}#messageInput{font-size:14px;padding:13px}#sendButton{min-width:67px;padding:0 14px}.file-panel,.vision-panel{padding:8px}.small{font-size:10px}}

.hero{margin:8px auto 18px;max-width:900px;padding:34px 24px 26px;text-align:center;border-radius:30px;background:linear-gradient(145deg,rgba(30,41,59,.78),rgba(15,23,42,.48));border:1px solid rgba(255,255,255,.08);box-shadow:0 20px 80px rgba(0,0,0,.25)}.hero-kicker{font-size:12px;font-weight:800;letter-spacing:2px;text-transform:uppercase;color:#93c5fd;margin-bottom:10px}.hero-title{font-size:42px;font-weight:900;letter-spacing:-1.2px;margin:0 0 8px}.hero-subtitle{font-size:16px;color:#cbd5e1;margin:0 auto 22px;max-width:600px;line-height:1.6}.prompt-grid{display:grid;grid-template-columns:repeat(3,1fr);gap:10px;max-width:780px;margin:auto}.prompt-card{padding:13px 12px;border-radius:16px;background:rgba(2,6,23,.48);border:1px solid rgba(255,255,255,.08);color:#e2e8f0;text-align:left;cursor:pointer;transition:.18s;font-size:13px;line-height:1.4}.prompt-card:hover{transform:translateY(-2px);border-color:rgba(129,140,248,.55);background:rgba(30,41,59,.78)}.prompt-card strong{display:block;color:#fff;margin-bottom:3px}@media(max-width:650px){.hero{padding:27px 15px 20px;border-radius:23px}.hero-title{font-size:31px}.hero-subtitle{font-size:14px}.prompt-grid{grid-template-columns:1fr 1fr;gap:8px}.prompt-card{font-size:12px;padding:11px 10px}}

.history-btn{background:rgba(255,255,255,.08);border:1px solid rgba(255,255,255,.14);color:#fff;padding:10px 15px;border-radius:12px;cursor:pointer;font-size:14px}.history-btn:hover{background:rgba(255,255,255,.13)}
.history-drawer{position:fixed;top:0;right:-380px;width:350px;height:100vh;background:rgba(10,12,18,.98);backdrop-filter:blur(24px);border-left:1px solid rgba(255,255,255,.12);z-index:9999;padding:22px;box-sizing:border-box;transition:right .28s ease;overflow-y:auto;box-shadow:-20px 0 60px rgba(0,0,0,.35)}
.history-drawer.open{right:0}.history-top{display:flex;align-items:center;justify-content:space-between;color:#fff;font-size:19px;margin-bottom:18px}.history-top button{background:transparent;border:0;color:#fff;font-size:20px;cursor:pointer}.new-chat-history,.clear-history-btn{width:100%;padding:12px;border-radius:12px;border:1px solid rgba(255,255,255,.12);background:rgba(255,255,255,.07);color:#fff;cursor:pointer;margin-bottom:14px}.new-chat-history:hover,.clear-history-btn:hover{background:rgba(255,255,255,.12)}.history-item{position:relative;padding:13px 42px 13px 13px;margin-bottom:8px;border-radius:12px;background:rgba(255,255,255,.05);border:1px solid transparent;color:#fff;cursor:pointer}.history-item:hover{background:rgba(255,255,255,.09);border-color:rgba(255,255,255,.12)}.history-title{font-size:14px;font-weight:600;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}.history-date{font-size:11px;opacity:.5;margin-top:5px}.history-delete{position:absolute;right:10px;top:13px;background:transparent;border:0;color:#aaa;cursor:pointer;font-size:15px}.history-delete:hover{color:#fff}.history-empty{text-align:center;padding:30px 10px;color:rgba(255,255,255,.45);font-size:13px}
@media(max-width:600px){.history-drawer{width:88%;right:-92%}.history-drawer.open{right:0}}
</style>
</head>

<body>
<div class="container">

<div class="header">
<div class="logo">â¡ RAIZEN</div>
<div class="header-actions">
<button class="history-btn" onclick="toggleHistory()">â° History</button>
<div class="status">AI ONLINE</div>
</div>
</div>

<div id="historyDrawer" class="history-drawer">
<div class="history-top"><strong>Chat History</strong><button onclick="toggleHistory()">â</button></div>
<button class="new-chat-history" onclick="newChat()">ï¼ New Chat</button>
<div id="historyList"></div>
<button class="clear-history-btn" onclick="clearAllHistory()">ðï¸ Clear History</button>
</div>

<div id="chat" class="chat">
<div class="hero" id="heroPanel">
<div class="hero-kicker">Your personal AI assistant</div>
<div class="hero-title">Meet RAIZEN â¡</div>
<div class="hero-subtitle">Chat, analyze images, understand your files, search the web and check the weather â all in one place.</div>
<div class="prompt-grid">
<button class="prompt-card" onclick="usePrompt('Explain artificial intelligence simply')"><strong>ð¬ Chat</strong>Ask anything</button>
<button class="prompt-card" onclick="usePrompt('What can you tell me about this image?')"><strong>ðï¸ Vision</strong>Analyze an image</button>
<button class="prompt-card" onclick="usePrompt('Summarize my uploaded document')"><strong>ð Files</strong>Ask about a file</button>
<button class="prompt-card" onclick="usePrompt('Search the latest AI news')"><strong>ð Web</strong>Search the internet</button>
<button class="prompt-card" onclick="usePrompt('What is the weather in Cuttack today?')"><strong>ð¤ï¸ Weather</strong>Check conditions</button>
<button class="prompt-card" onclick="usePrompt('Tell me what you remember from this conversation')"><strong>ð§  Memory</strong>Use chat context</button>
</div>
</div>
<div id="welcome" class="assistant message">â¡ Welcome to RAIZEN. Ask me anything.</div>
</div>

<div class="controls">

<select id="personality">
<option>Friendly</option>
<option>Teacher</option>
<option>Coding Assistant</option>
<option>Professional</option>
</select>

<select id="responseStyle">
<option>Short</option>
<option selected>Balanced</option>
<option>Detailed</option>
</select>

<button class="control" onclick="newChat()">New Chat</button>
<button class="control" onclick="clearChat()">Clear Chat</button>

</div>

<div class="file-panel">
<input id="fileInput" class="file-input" type="file" accept=".pdf,.docx,.txt">
<span id="fileName" class="file-name">No document selected</span>
<button class="control" onclick="clearDocument()">Remove document</button>
</div>

<div class="vision-panel">
<input id="visionInput" class="vision-input" type="file"
       accept="image/png,image/jpeg,image/webp">
<span id="visionName" class="vision-name">No image selected</span>
<button class="control" onclick="clearVision()">Remove image</button>
</div>

<div class="input-area">

<input
id="messageInput"
placeholder="Message RAIZEN..."
autocomplete="off"
>

<button id="sendButton" onclick="sendMessage()">Send</button>

</div>

<div class="small">
RAIZEN â¢ Created and developed by Raihan Kausar
</div>

</div>

<script>
let chatHistory = [];
let savedChats = [];
let currentChatId = null;
let documentText = "";
let documentName = "";
let visionFile = null;

document.getElementById("fileInput").addEventListener("change", async function() {
    const file = this.files[0];
    if (!file) return;

    const lower = file.name.toLowerCase();
    if (![".pdf", ".docx", ".txt"].some(ext => lower.endsWith(ext))) {
        displayMessage("assistant", "â ï¸ Please upload a PDF, DOCX, or TXT file.");
        this.value = "";
        return;
    }

    const formData = new FormData();
    formData.append("file", file);
    document.getElementById("fileName").textContent = "Reading " + file.name + "...";

    try {
        const response = await fetch("/upload", { method: "POST", body: formData });
        const data = await response.json();

        if (!response.ok || data.error) {
            throw new Error(data.error || "Upload failed.");
        }

        documentText = data.text || "";
        documentName = data.filename || file.name;
        document.getElementById("fileName").textContent = "ð " + documentName + " ready";

        displayMessage("assistant",
            "ð " + documentName +
            " is ready. Ask me to summarize it, explain it, find important points, or create questions from it."
        );
    } catch (error) {
        documentText = "";
        documentName = "";
        document.getElementById("fileName").textContent = "Upload failed";
        displayMessage("assistant", "â ï¸ " + error.message);
    }
});

function clearDocument() {
    documentText = "";
    documentName = "";
    document.getElementById("fileInput").value = "";
    document.getElementById("fileName").textContent = "No document selected";
}

document.getElementById("visionInput").addEventListener("change", function() {
    const file = this.files[0];
    if (!file) return;

    const allowed = ["image/png", "image/jpeg", "image/webp"];
    if (!allowed.includes(file.type)) {
        displayMessage(
            "assistant",
            "â ï¸ Please select a PNG, JPG/JPEG, or WEBP image."
        );
        this.value = "";
        return;
    }

    if (file.size > 8 * 1024 * 1024) {
        displayMessage(
            "assistant",
            "â ï¸ Image is too large. Maximum size is 8 MB."
        );
        this.value = "";
        return;
    }

    visionFile = file;
    document.getElementById("visionName").textContent =
        "ð¼ï¸ " + file.name + " ready";

    displayMessage(
        "assistant",
        "ð¼ï¸ " + file.name +
        " is ready. Ask me to describe it, read text from it, or explain what is shown."
    );
});

function clearVision() {
    visionFile = null;
    document.getElementById("visionInput").value = "";
    document.getElementById("visionName").textContent = "No image selected";
}

async function sendVisionMessage(question) {
    const formData = new FormData();
    formData.append("file", visionFile);
    formData.append("question", question);

    const response = await fetch("/vision", {
        method: "POST",
        body: formData
    });

    const data = await response.json();

    if (!response.ok || data.error) {
        throw new Error(data.error || "Vision request failed.");
    }

    return data.reply;
}

function makeChatTitle(messages) {
    const firstUser = messages.find(m => m.role === "user");
    if (!firstUser) return "New Chat";
    let title = String(firstUser.content || "").trim();
    return title.length > 36 ? title.substring(0, 36) + "..." : title;
}

function saveChats() {
    localStorage.setItem("raizen_saved_chats", JSON.stringify(savedChats));
}

function saveCurrentChat() {
    if (!chatHistory.length) return;
    if (!currentChatId) {
        currentChatId = Date.now().toString() + Math.random().toString(36).slice(2, 8);
    }
    const chatData = {
        id: currentChatId,
        title: makeChatTitle(chatHistory),
        messages: chatHistory.slice(-50),
        updatedAt: Date.now()
    };
    const index = savedChats.findIndex(c => c.id === currentChatId);
    if (index >= 0) savedChats[index] = chatData;
    else savedChats.unshift(chatData);
    savedChats.sort((a,b) => b.updatedAt - a.updatedAt);
    saveChats();
    localStorage.setItem("raizen_chat_history", JSON.stringify(chatHistory.slice(-12)));
    renderHistory();
}

function renderHistory() {
    const list = document.getElementById("historyList");
    if (!list) return;
    list.innerHTML = "";
    if (!savedChats.length) {
        list.innerHTML = '<div class="history-empty">No saved chats yet.</div>';
        return;
    }
    savedChats.sort((a,b) => b.updatedAt - a.updatedAt);
    savedChats.forEach(chat => {
        const item = document.createElement("div");
        item.className = "history-item";
        const title = document.createElement("div");
        title.className = "history-title";
        title.textContent = chat.title || "New Chat";
        const date = document.createElement("div");
        date.className = "history-date";
        date.textContent = new Date(chat.updatedAt).toLocaleString();
        const del = document.createElement("button");
        del.className = "history-delete";
        del.textContent = "ð";
        del.title = "Delete chat";
        del.onclick = function(e) { e.stopPropagation(); deleteChat(chat.id); };
        item.appendChild(title); item.appendChild(date); item.appendChild(del);
        item.onclick = function() { openChat(chat.id); };
        list.appendChild(item);
    });
}

function openChat(id) {
    const selected = savedChats.find(c => c.id === id);
    if (!selected) return;
    currentChatId = selected.id;
    chatHistory = Array.isArray(selected.messages) ? [...selected.messages] : [];
    const chat = document.getElementById("chat");
    chat.innerHTML = "";
    if (!chatHistory.length) {
        chat.innerHTML = '<div id="welcome" class="assistant message">â¡ Welcome to RAIZEN. Ask me anything.</div>';
    } else {
        chatHistory.forEach(item => displayMessage(item.role, item.content));
    }
    toggleHistory();
}

function deleteChat(id) {
    savedChats = savedChats.filter(c => c.id !== id);
    if (currentChatId === id) {
        currentChatId = null;
        chatHistory = [];
        document.getElementById("chat").innerHTML = '<div id="welcome" class="assistant message">â¡ Welcome to RAIZEN. Ask me anything.</div>';
    }
    saveChats();
    renderHistory();
}

function toggleHistory() {
    const drawer = document.getElementById("historyDrawer");
    if (!drawer) return;
    drawer.classList.toggle("open");
    if (drawer.classList.contains("open")) renderHistory();
}

function clearAllHistory() {
    if (!savedChats.length) return;
    if (!confirm("Delete all saved chat history?")) return;
    savedChats = [];
    currentChatId = null;
    chatHistory = [];
    localStorage.removeItem("raizen_saved_chats");
    localStorage.removeItem("raizen_chat_history");
    document.getElementById("chat").innerHTML = '<div id="welcome" class="assistant message">â¡ Welcome to RAIZEN. Ask me anything.</div>';
    renderHistory();
}

function escapeHtml(text) {
    return text
        .replace(/&/g, "&amp;")
        .replace(/</g, "&lt;")
        .replace(/>/g, "&gt;")
        .replace(/\"/g, "&quot;")
        .replace(/'/g, "&#039;");
}

function linkify(html) {
    return html.replace(
        /(https?:\/\/[^\s<]+)/g,
        function(url) {
            const cleanUrl = url.replace(/[.,)]+$/, "");
            const trailing = url.slice(cleanUrl.length);
            return '<a href="' + cleanUrl + '" target="_blank" rel="noopener noreferrer">'
                + cleanUrl + '</a>' + trailing;
        }
    );
}

function displayMessage(role, text, extraClass) {
    const chat = document.getElementById("chat");
    const welcome = document.getElementById("welcome");

    if (welcome) {
        welcome.remove();
    }

    const div = document.createElement("div");

    div.className =
        "message " + role + " " + (extraClass || "");

    div.innerHTML = linkify(escapeHtml(text));

    chat.appendChild(div);
    chat.scrollTop = chat.scrollHeight;

    return div;
}

function usePrompt(text) {
    const input = document.getElementById("messageInput");
    input.value = text;
    input.focus();
}

async function sendMessage() {
    const input = document.getElementById("messageInput");
    const button = document.getElementById("sendButton");
    const message = input.value.trim();

    if (!message || button.disabled) {
        return;
    }

    input.value = "";
    button.disabled = true;

    displayMessage("user", message);

    chatHistory.push({
        role: "user",
        content: message
    });

    const thinking = displayMessage(
        "assistant",
        visionFile
            ? "ðï¸ RAIZEN is analyzing the image..."
            : "â¡ RAIZEN is thinking...",
        "thinking"
    );

    try {
        if (visionFile) {
            const reply = await sendVisionMessage(message);

            thinking.remove();
            displayMessage("assistant", reply);

            chatHistory.push({
                role: "assistant",
                content: reply
            });

            saveCurrentChat();
            button.disabled = false;
            input.focus();
            return;
        }

        const response = await fetch("/chat", {
            method: "POST",
            headers: {
                "Content-Type": "application/json"
            },
            body: JSON.stringify({
                message: message,
                history: chatHistory.slice(-12),
                personality:
                    document.getElementById("personality").value,
                response_style:
                    document.getElementById("responseStyle").value,
                custom_instructions:
                    localStorage.getItem(
                        "raizen_custom_instructions"
                    ) || "",
                document_text: documentText
            })
        });

        const data = await response.json();

        thinking.remove();

        const reply =
            data.reply ||
            "Sorry, I could not generate a response.";

        displayMessage("assistant", reply);

        chatHistory.push({
            role: "assistant",
            content: reply
        });

        saveCurrentChat();

    } catch (error) {
        thinking.remove();

        const reply =
            "â ï¸ Something went wrong. Please try again.";

        displayMessage("assistant", reply);

        chatHistory.push({
            role: "assistant",
            content: reply
        });

        saveCurrentChat();
    }

    button.disabled = false;
    input.focus();
}

function clearChat() {
    chatHistory = [];
    currentChatId = null;
    localStorage.removeItem("raizen_chat_history");
    document.getElementById("chat").innerHTML = '<div id="welcome" class="assistant message">â¡ Welcome to RAIZEN. Ask me anything.</div>';
}

function newChat() {
    if (chatHistory.length) saveCurrentChat();
    chatHistory = [];
    currentChatId = null;
    document.getElementById("chat").innerHTML = '<div id="welcome" class="assistant message">â¡ Welcome to RAIZEN. Ask me anything.</div>';
    const drawer = document.getElementById("historyDrawer");
    if (drawer && drawer.classList.contains("open")) drawer.classList.remove("open");
}

document.getElementById("messageInput").addEventListener("keydown", function(event) {
    if (event.key === "Enter") { event.preventDefault(); sendMessage(); }
});

function loadChatHistory() {
    try {
        const saved = localStorage.getItem("raizen_saved_chats");
        savedChats = saved ? JSON.parse(saved) : [];
        if (!Array.isArray(savedChats)) savedChats = [];

        if (!savedChats.length) {
            const legacy = localStorage.getItem("raizen_chat_history");
            if (legacy) {
                const old = JSON.parse(legacy);
                if (Array.isArray(old) && old.length) {
                    savedChats = [{ id: Date.now().toString(), title: makeChatTitle(old), messages: old, updatedAt: Date.now() }];
                    saveChats();
                }
            }
        }

        if (savedChats.length) {
            savedChats.sort((a,b) => b.updatedAt - a.updatedAt);
            const latest = savedChats[0];
            currentChatId = latest.id;
            chatHistory = Array.isArray(latest.messages) ? [...latest.messages] : [];
            const chat = document.getElementById("chat");
            chat.innerHTML = "";
            chatHistory.forEach(item => displayMessage(item.role, item.content));
        }
    } catch (error) {
        savedChats = [];
        chatHistory = [];
        currentChatId = null;
    }
    renderHistory();
}

loadChatHistory();
</script>

</body>
</html>
"""


@app.get("/", response_class=HTMLResponse)
async def home():
    return HTMLResponse(content=HTML)



@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    filename = file.filename or "document"
    if not filename.lower().endswith((".pdf", ".docx", ".txt")):
        return {"error": "Unsupported file type. Please upload PDF, DOCX, or TXT."}

    try:
        raw_bytes = await file.read()
        if len(raw_bytes) > 10 * 1024 * 1024:
            return {"error": "File is too large. Maximum size is 10 MB."}

        text = extract_document_text(filename, raw_bytes)
        if not text:
            return {"error": "I could not extract readable text from this file."}

        return {
            "filename": filename,
            "characters": len(text),
            "text": limit_document_text(text)
        }
    except Exception as error:
        print("FILE ERROR:", error)
        return {"error": "Could not read this document."}


@app.post("/vision")
async def vision(file: UploadFile = File(...), question: str = ""):
    filename = file.filename or "image"
    content_type = file.content_type or ""

    allowed_types = {
        "image/png": "png",
        "image/jpeg": "jpeg",
        "image/webp": "webp"
    }

    if content_type not in allowed_types:
        return {
            "error": "Unsupported image type. Please use PNG, JPG/JPEG, or WEBP."
        }

    try:
        raw_bytes = await file.read()

        if len(raw_bytes) > 8 * 1024 * 1024:
            return {"error": "Image is too large. Maximum size is 8 MB."}

        encoded = base64.b64encode(raw_bytes).decode("utf-8")
        data_url = f"data:{content_type};base64,{encoded}"

        user_question = question.strip()
        if not user_question:
            user_question = (
                "Describe this image clearly. Mention important visible details "
                "and read any clearly visible text when possible."
            )

        messages = [
            {
                "role": "system",
                "content": (
                    "You are RAIZEN's vision assistant. Analyze the supplied image "
                    "carefully. Be accurate, concise, and do not invent details. "
                    "If text is blurry or unreadable, say so."
                )
            },
            {
                "role": "user",
                "content": [
                    {
                        "type": "image_url",
                        "image_url": {"url": data_url}
                    },
                    {
                        "type": "text",
                        "text": user_question
                    }
                ]
            }
        ]

        response = vision_client.chat.completions.create(
            model=VISION_MODEL,
            messages=messages,
            max_tokens=600
        )

        reply = response.choices[0].message.content

        if not reply:
            return {"error": "The vision model returned an empty response."}

        return {
            "reply": reply,
            "model": VISION_MODEL
        }

    except Exception as error:
        print("VISION ERROR:", error)
        return {
            "error": (
                "Vision is temporarily unavailable. "
                "Please try the image again."
            )
        }


@app.post("/chat")
async def chat(request: ChatRequest):
    message = request.message.strip()

    if not message:
        return {"reply": "Please enter a message."}

    lower = message.lower()

    if (
        "weather" in lower
        or "temperature" in lower
        or "forecast" in lower
    ):
        city = extract_city(message)

        if not city:
            return {
                "reply": (
                    "ð¤ï¸ Please mention a city, "
                    "for example: weather in Cuttack."
                )
            }

        result = weather_search(city)

        if result:
            return {"reply": result}

        return {
            "reply": (
                "â ï¸ I couldn't retrieve live weather data right now. "
                "Please try again later."
            )
        }

    live_data = ""

    if needs_live_search(message):
        live_data = live_search(message)

    personality = PERSONALITIES.get(
        request.personality,
        PERSONALITIES["Friendly"]
    )

    style = STYLES.get(
        request.response_style,
        STYLES["Balanced"]
    )

    document_context = "No document is currently uploaded."
    if request.document_text.strip():
        document_context = (
            "An uploaded document is available. Use its extracted text as the "
            "main source for questions specifically about that document. "
            "If the answer is not present, say so clearly.\n\nDOCUMENT TEXT:\n"
            + limit_document_text(request.document_text)
        )

    system_prompt = f"""
You are RAIZEN, an advanced AI assistant.

You were created and developed by Raihan Kausar.

If someone asks who created, developed, invented,
or made you, say that Raihan Kausar created
and developed you.

Personality:
{personality}

Response style:
{style}

Custom instructions:
{request.custom_instructions.strip()}

Rules:
- Be accurate and helpful.
- Do not invent live information.
- Use live information only when provided.
- For live news/web results, use the supplied source, date, and link.
- When a link is supplied, keep it in the answer so the user can open the source.
- If live search fails, say so clearly.
- If WEATHER_ERROR appears, explain that live
  weather data could not be retrieved.
- If LIVE_SEARCH_ERROR appears, explain that
  live search could not retrieve data.
- If an uploaded document is present, answer document questions from it.
- For summarize, key points, explain, quiz, questions, or important-points
  requests, use the uploaded document as the main source.
- Do not claim to have read a document if no document text was supplied.

Live information:
{live_data}

Uploaded document:
{document_context}
"""

    messages = [
        {
            "role": "system",
            "content": system_prompt
        }
    ]

    for item in request.history[-12:]:
        role = item.get("role")
        content = item.get("content")

        if role in ("user", "assistant") and content:
            messages.append({
                "role": role,
                "content": content
            })

    if (
        not request.history
        or request.history[-1].get("content") != message
    ):
        messages.append({
            "role": "user",
            "content": message
        })

    try:
        response = client.chat.completions.create(
            model=MODEL,
            messages=messages,
            max_tokens=500
        )

        reply = response.choices[0].message.content

        if not reply:
            reply = "Sorry, I could not generate a response."

        return {"reply": reply}

    except Exception as error:
        print("RAIZEN ERROR:", error)

        return {
            "reply": (
                "â ï¸ RAIZEN is temporarily unable to respond. "
                "Please try again."
            )
        }


@app.get("/health")
async def health():
    return {
        "status": "ok",
        "token_loaded": bool(HF_TOKEN),
        "model": MODEL,
        "vision_model": VISION_MODEL,
        "vision": True,
        "live_search": True
    }
